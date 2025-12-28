from __future__ import annotations

from pathlib import Path
import time
import uuid
import logging
import hashlib
import threading
from concurrent.futures import ThreadPoolExecutor
from typing import Optional, Dict, List, Union, Tuple, Any, cast
import json

from fastapi import FastAPI, HTTPException, UploadFile, File, Header, Depends, Form, Body
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from starlette.responses import PlainTextResponse

from .agents import AgentOrchestrator
from .config import get_settings
from .models import (
    HealthStatus,
    QueryRequest,
    QueryResponse,
    SessionCreateRequest,
    SessionCreateResponse,
    SessionHistoryResponse,
    SessionListResponse,
    SessionSummary,
    SessionUpdateRequest,
)
from .vectorstore import VectorStore
from .rag import load_documents
from .document_storage import DocumentStorage
import pypdf
import docx
import openpyxl
from .session_store import SessionStore
from .metrics import Metrics
from .user_store import UserStore
from .redis_coord import RedisCoordinator

logger = logging.getLogger("campusqa.main")
if not logger.handlers:
    logging.basicConfig(level=logging.INFO)

settings = get_settings()
orchestrator = AgentOrchestrator(settings)
vectorstore = VectorStore(settings)
session_store = SessionStore(settings)
metrics = Metrics()
user_store = UserStore(settings)
doc_storage = DocumentStorage(settings)
redis_coord = RedisCoordinator(settings)
ingest_executor = ThreadPoolExecutor(max_workers=max(1, settings.ingest_workers))
_vectorize_lock = threading.Lock()

if settings.use_oss_storage:
    try:
        doc_storage.sync_from_remote()
    except Exception as exc:
        logger.warning("Initial OSS sync failed: %s", exc)

app = FastAPI(title=settings.project_name)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

static_dir = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=static_dir), name="static")

ALLOWED_DOC_EXTS = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


def _extract_text_for_preview(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        with path.open("rb") as f:
            reader = pypdf.PdfReader(f)
            texts: List[str] = []
            for page in reader.pages:
                try:
                    texts.append(page.extract_text() or "")
                except Exception:
                    continue
            return "\n".join(texts)
    if ext == ".docx":
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts: List[str] = []
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" \t ".join(cells))
        wb.close()
        return "\n".join(parts)
    return ""


def _parse_bearer_token(authorization: Optional[str]) -> str:
    if not authorization:
        raise HTTPException(status_code=401, detail="缺少凭证")
    parts = authorization.split()
    if len(parts) != 2 or parts[0].lower() != "bearer":
        raise HTTPException(status_code=401, detail="凭证格式错误")
    return parts[1]


def _require_admin(authorization: Optional[str] = Header(default=None)) -> Tuple[str, str]:
    """管理员只校验 token & 角色，不强制首次改密。"""

    token = _parse_bearer_token(authorization)
    auth = user_store.validate_token(token)
    if not auth:
        raise HTTPException(status_code=401, detail="凭证无效或已过期")
    username, role, _ = auth
    if role != "admin":
        raise HTTPException(status_code=403, detail="需要管理员权限")
    return username, role


def _require_user_active(authorization: Optional[str] = Header(default=None)) -> str:
    """普通用户访问受保护资源，需要已修改初始密码。"""

    token = _parse_bearer_token(authorization)
    auth = user_store.validate_token(token)
    if not auth:
        raise HTTPException(status_code=401, detail="凭证无效或已过期")
    username, _, must_change = auth
    if must_change:
        raise HTTPException(status_code=403, detail="请先修改初始密码")
    return username
def _safe_doc_path(name: str) -> Path:
    if "/" in name or ".." in name:
        raise HTTPException(status_code=400, detail="非法文件名")
    try:
        path = doc_storage.ensure_local(name)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="文件不存在")
    if path.suffix.lower() not in ALLOWED_DOC_EXTS:
        raise HTTPException(status_code=400, detail="不支持的文件类型")
    return path



def _require_user_allow_change(authorization: Optional[str] = Header(default=None)) -> str:
    """仅用于修改密码，允许仍在 must_change 状态的 token。"""

    token = _parse_bearer_token(authorization)
    auth = user_store.validate_token(token)
    if not auth:
        raise HTTPException(status_code=401, detail="凭证无效或已过期")
    username, _, _ = auth
    return username


def _record_ingest_event(doc_name: str, status: str, **meta: Any) -> None:
    payload: Dict[str, Any] = {"doc": doc_name, "status": status}
    if meta:
        payload.update(meta)
    redis_coord.record_event(payload)


def _run_vectorize_job(entries: List[Dict[str, Any]]) -> None:
    filenames = [cast(str, entry.get("filename")) for entry in entries if entry.get("filename")]
    if not filenames:
        return
    job_map = {cast(str, entry.get("filename")): entry.get("job_id") for entry in entries if entry.get("filename")}
    try:
        with _vectorize_lock:
            chunks = load_documents(
                settings.docs_path,
                settings.chunk_size,
                settings.chunk_overlap,
                settings.chunk_overlap_ratio,
            )
            vectorstore.build(chunks)
            docs_count = len(chunks)
    except Exception as exc:  # pragma: no cover - background job should not crash server
        logger.exception("后台索引构建失败: %s", exc)
        for name in filenames:
            redis_coord.set_status(name, "failed", {"job_id": job_map.get(name), "error": "索引构建失败"})
            _record_ingest_event(name, "failed", job_id=job_map.get(name), error="索引构建失败")
        return

    for name in filenames:
        metrics.record_doc_upload()
        redis_coord.set_status(name, "completed", {"chunk_count": docs_count, "job_id": job_map.get(name)})
        _record_ingest_event(name, "completed", job_id=job_map.get(name), chunk_count=docs_count)


def _schedule_vectorize(entries: List[Dict[str, Any]]) -> None:
    if not entries:
        return
    payload = [dict(item) for item in entries]
    try:
        ingest_executor.submit(_run_vectorize_job, payload)
    except Exception as exc:  # pragma: no cover - fallback to inline for reliability
        logger.warning("提交后台索引任务失败，改为同步执行: %s", exc)
        _run_vectorize_job(payload)


@app.get("/health", response_model=HealthStatus)
def health() -> HealthStatus:
    try:
        docs_indexed, _ = vectorstore.stats()
    except FileNotFoundError:
        docs_indexed = 0
    return HealthStatus(status="ok", embedding_model=settings.embedding_model, docs_indexed=docs_indexed)


@app.post("/api/query", response_model=QueryResponse)
def query(req: QueryRequest, user: str = Depends(_require_user_active)) -> QueryResponse:
    start = time.perf_counter()
    req_id = uuid.uuid4().hex[:8]
    from_cache = False
    logger.info("[query:%s] start user=%s session=%s top_k=%s", req_id, user, req.session_id, req.top_k)
    try:
        resp, from_cache = orchestrator.handle(req)
        # persist conversation if session_id provided
        if req.session_id:
            session_store.ensure_session_for_user(req.session_id, user, create_if_missing=True)
            session_store.add_message(req.session_id, "user", req.query, user=user)
            session_store.add_message(req.session_id, "assistant", resp.answer, user=user)
        latency_ms = (time.perf_counter() - start) * 1000
        metrics.record_query(latency_ms=latency_ms, cached=from_cache, used_tools=len(resp.used_tools))
        logger.info(
            "[query:%s] done latency_ms=%.1f cached=%s sources=%s tools=%s",
            req_id,
            latency_ms,
            from_cache,
            len(resp.sources) if hasattr(resp, "sources") else "-",
            len(resp.used_tools) if hasattr(resp, "used_tools") else "-",
        )
        return resp
    except FileNotFoundError as e:
        metrics.record_error()
        logger.exception("[query:%s] failed: %s", req_id, e)
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/query/stream")
def query_stream(req: QueryRequest, user: str = Depends(_require_user_active)):
    start = time.perf_counter()
    req_id = uuid.uuid4().hex[:8]
    logger.info("[query-stream:%s] start user=%s session=%s top_k=%s", req_id, user, req.session_id, req.top_k)
    try:
        gen = orchestrator.handle_stream(req)
        answer_parts: List[str] = []
        meta_payload: Dict[str, Any] = {}

        def wrapped():
            nonlocal meta_payload
            try:
                for chunk in gen:
                    if chunk.startswith("__META__"):
                        try:
                            meta_payload = json.loads(chunk[len("__META__"):])
                        except Exception:
                            logger.warning("[query-stream:%s] failed to parse meta chunk", req_id)
                        yield chunk
                        continue
                    answer_parts.append(chunk)
                    yield chunk
            finally:
                latency_ms = (time.perf_counter() - start) * 1000
                metrics.record_stream(latency_ms=latency_ms, count=False)
                logger.info("[query-stream:%s] done latency_ms=%.1f", req_id, latency_ms)
                _persist_stream_history(req, user, answer_parts, meta_payload)

        # record start for streaming count even if client drops
        metrics.record_stream()
        return StreamingResponse(wrapped(), media_type="text/plain")
    except FileNotFoundError as e:
        metrics.record_error()
        logger.exception("[query-stream:%s] failed: %s", req_id, e)
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/admin/login")
def admin_login(username: str = Form(...), password: str = Form(...)) -> Dict[str, Union[str, bool]]:
    ok, token, must_change, role = user_store.authenticate(username, password)
    if not ok or not token or role != "admin":
        raise HTTPException(status_code=401, detail="用户名或密码错误，或无管理员权限")
    return {"token": token, "must_change_password": must_change, "role": role}
@app.post("/api/admin/users/batch_register")
def batch_register(student_ids: List[str] = Body(...), admin: Tuple[str, str] = Depends(_require_admin)) -> Dict[str, Union[List[str], str]]:
    if not student_ids:
        raise HTTPException(status_code=400, detail="student_ids 不能为空")
    created, skipped = user_store.batch_create(student_ids)
    return {"created": created, "skipped": skipped, "initial_password_rule": "hziee+学号"}


@app.post("/api/auth/login")
def user_login(username: str = Form(...), password: str = Form(...)) -> Dict[str, Union[str, bool]]:
    ok, token, must_change, role = user_store.authenticate(username, password)
    if not ok or not token:
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    return {"token": token, "must_change_password": must_change, "role": role or "student"}


@app.post("/api/auth/change_password")
def change_password(new_password: str = Form(...), user: str = Depends(_require_user_allow_change)):
    if not new_password or len(new_password) < 6:
        raise HTTPException(status_code=400, detail="新密码至少6位")
    user_store.set_password(user, new_password)
    return {"status": "ok"}


@app.post("/api/docs/upload")
def upload_docs(files: List[UploadFile] = File(...), admin: Tuple[str, str] = Depends(_require_admin)):
    if not files:
        raise HTTPException(status_code=400, detail="请至少上传一个文件")

    admin_user, _ = admin
    successes: List[Dict[str, Any]] = []
    failures: List[Dict[str, Any]] = []
    existing_docs = doc_storage.list_documents()
    existing_hashes: Dict[str, Dict[str, Any]] = {
        cast(str, entry["hash"]): entry for entry in existing_docs if entry.get("hash")
    }
    batch_hashes: set[str] = set()
    max_bytes = max(0, settings.max_upload_mb) * 1024 * 1024 if settings.max_upload_mb else 0

    for upload in files:
        filename = (upload.filename or "").strip()
        if not filename:
            failures.append({"filename": "(未命名)", "reason": "文件名为空"})
            continue
        filename = Path(filename).name  # 防止路径穿越
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_DOC_EXTS:
            failures.append({"filename": filename, "reason": f"不支持的文件类型: {ext}"})
            continue

        content = upload.file.read()
        if max_bytes and len(content) > max_bytes:
            failures.append({"filename": filename, "reason": f"文件超过大小限制 {settings.max_upload_mb}MB"})
            continue

        file_hash = doc_storage.compute_hash_from_bytes(content)
        if file_hash in batch_hashes:
            failures.append({"filename": filename, "reason": "与本次上传的其他文件内容相同，已跳过"})
            continue

        existing_entry = existing_hashes.get(file_hash)
        if existing_entry:
            existing_name = existing_entry.get("name") or existing_entry.get("filename") or "已存在文件"
            failures.append({"filename": filename, "reason": f"内容与 {existing_name} 重复"})
            continue

        batch_hashes.add(file_hash)

        job_id = uuid.uuid4().hex
        redis_coord.enqueue(
            "doc_ingest",
            {"job_id": job_id, "filename": filename, "admin": admin_user, "size": len(content), "ts": time.time()},
        )
        redis_coord.set_status(filename, "pending", {"job_id": job_id})
        _record_ingest_event(filename, "pending", job_id=job_id, admin=admin_user, size=len(content), hash=file_hash)

        coord_enabled = redis_coord.enabled
        with redis_coord.lock(f"doc-ingest:{filename}", ttl=600, wait_timeout=5) as lock_handle:
            if coord_enabled and lock_handle is None:
                redis_coord.set_status(filename, "busy", {"job_id": job_id})
                failures.append({"filename": filename, "reason": "文档正在处理中"})
                _record_ingest_event(filename, "busy", job_id=job_id, reason="locked")
                continue
            try:
                redis_coord.set_status(filename, "uploading", {"size": len(content)})
                _record_ingest_event(filename, "uploading", job_id=job_id, size=len(content))
                doc_storage.save(filename, content)

                if settings.use_oss_storage:
                    try:
                        doc_storage.sync_from_remote()
                    except Exception as exc:
                        logger.warning("同步 OSS 失败: %s", exc)

                redis_coord.set_status(filename, "vectorizing", {"job_id": job_id})
                _record_ingest_event(filename, "vectorizing", job_id=job_id)
                successes.append({"filename": filename, "job_id": job_id, "hash": file_hash})
                existing_hashes[file_hash] = {"name": filename}
            except HTTPException as exc:
                redis_coord.set_status(filename, "failed", {"job_id": job_id, "error": exc.detail})
                failures.append({"filename": filename, "reason": exc.detail})
                _record_ingest_event(filename, "failed", job_id=job_id, error=exc.detail)
            except Exception as exc:
                redis_coord.set_status(filename, "failed", {"job_id": job_id, "error": str(exc)})
                failures.append({"filename": filename, "reason": "处理失败"})
                logger.exception("上传文件失败: %s", exc)
                _record_ingest_event(filename, "failed", job_id=job_id, error=str(exc))

    vectorization_info: Dict[str, Any] = {"mode": "idle", "scheduled": False, "pending": []}
    if successes:
        vectorization_info = {
            "mode": "async",
            "scheduled": True,
            "pending": [entry["filename"] for entry in successes],
        }
        _schedule_vectorize(successes)
    try:
        docs_count, _ = vectorstore.stats()
    except FileNotFoundError:
        docs_count = 0

    overall_status = "ok"
    if failures and successes:
        overall_status = "partial"
    elif failures and not successes:
        overall_status = "failed"

    return JSONResponse(
        {
            "status": overall_status,
            "processed": [entry["filename"] for entry in successes],
            "failed": failures,
            "docs_count": docs_count,
            "vectorization": vectorization_info,
        }
    )


@app.delete("/api/admin/docs/{name}")
def delete_doc(name: str, admin: Tuple[str, str] = Depends(_require_admin)) -> Dict[str, Any]:
    try:
        doc_storage.delete(name)
        if settings.use_oss_storage:
            try:
                doc_storage.sync_from_remote()
            except Exception as exc:
                logger.warning("同步 OSS 失败: %s", exc)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="文件不存在")
    except Exception as exc:
        logger.exception("删除文档失败: %s", exc)
        raise HTTPException(status_code=500, detail="删除失败")

    chunks = load_documents(settings.docs_path, settings.chunk_size, settings.chunk_overlap, settings.chunk_overlap_ratio)
    vectorstore.build(chunks)
    return {"status": "deleted", "name": name, "docs_count": len(chunks)}


@app.post("/api/admin/docs/{name}/reindex")
def reindex_doc(name: str, admin: Tuple[str, str] = Depends(_require_admin)) -> Dict[str, Any]:
    admin_user, _ = admin
    path = _safe_doc_path(name)
    job_id = uuid.uuid4().hex
    entry: Dict[str, Any] = {"filename": path.name, "job_id": job_id, "admin": admin_user, "manual": True}
    redis_coord.set_status(path.name, "vectorizing", {"job_id": job_id, "note": "手动重建"})
    _record_ingest_event(path.name, "vectorizing", job_id=job_id, admin=admin_user, note="manual_reindex")
    _schedule_vectorize([entry])
    return {"status": "scheduled", "job_id": job_id}


@app.get("/api/admin/docs")
def list_docs(admin: Tuple[str, str] = Depends(_require_admin)) -> Dict[str, List[Dict[str, Any]]]:
    items: List[Dict[str, Any]] = []
    statuses = redis_coord.list_statuses() if redis_coord.enabled else {}
    for entry in doc_storage.list_documents():
        ext = str(entry.get("ext") or "").lower()
        if ext and ext not in ALLOWED_DOC_EXTS:
            continue
        name = str(entry.get("name") or "")
        status_meta = statuses.get(name, {})
        entry_payload: Dict[str, Any] = dict(entry)
        entry_payload.update({"status": status_meta.get("status", "unknown"), "status_meta": status_meta})
        items.append(entry_payload)
    return {"docs": items}


@app.get("/api/admin/docs/{name}", response_class=PlainTextResponse)
def get_doc_content(name: str, admin: Tuple[str, str] = Depends(_require_admin)):
    path = _safe_doc_path(name)
    content = _extract_text_for_preview(path)
    if not content:
        raise HTTPException(status_code=400, detail="无法读取文件内容")
    return content


@app.get("/api/docs/{name}", response_class=PlainTextResponse)
def get_doc_content_user(name: str, user: str = Depends(_require_user_active)):
    path = _safe_doc_path(name)
    content = _extract_text_for_preview(path)
    if not content:
        raise HTTPException(status_code=400, detail="无法读取文件内容")
    return content


@app.get("/api/docs/{name}/download")
def download_doc(name: str, user: str = Depends(_require_user_active)):
    path = _safe_doc_path(name)
    return FileResponse(path, filename=path.name)


@app.get("/api/admin/overview")
def admin_overview(admin: Tuple[str, str] = Depends(_require_admin)) -> Dict[str, Any]:
    metrics_snapshot = metrics.snapshot()
    try:
        docs_indexed, vectors_indexed = vectorstore.stats()
    except FileNotFoundError:
        docs_indexed, vectors_indexed = 0, 0

    statuses = redis_coord.list_statuses() if redis_coord.enabled else {}
    docs: List[Dict[str, Any]] = []
    for entry in doc_storage.list_documents():
        ext = str(entry.get("ext") or "").lower()
        if ext and ext not in ALLOWED_DOC_EXTS:
            continue
        name = str(entry.get("name") or "")
        status_meta = statuses.get(name, {})
        entry_payload: Dict[str, Any] = dict(entry)
        entry_payload.update({"status": status_meta.get("status", "unknown"), "status_meta": status_meta})
        docs.append(entry_payload)

    redis_snapshot = redis_coord.snapshot()
    recent_statuses: List[Dict[str, Any]] = []
    raw_statuses = redis_snapshot.get("statuses")
    typed_statuses: Dict[str, Dict[str, Any]] = {}
    if isinstance(raw_statuses, dict):
        for doc_name, payload in cast(Dict[Any, Any], raw_statuses).items():
            if not isinstance(payload, dict):
                continue
            payload_dict = dict(cast(Dict[str, Any], payload))
            typed_statuses[str(doc_name)] = payload_dict
    for doc_name, payload in typed_statuses.items():
        payload.setdefault("doc", doc_name)
        recent_statuses.append(payload)
    recent_statuses.sort(key=lambda item: float(item.get("ts", 0)), reverse=True)
    history = redis_snapshot.get("history")
    recent_history: List[Dict[str, Any]] = []
    if isinstance(history, list):
        for item in history:
            if isinstance(item, dict):
                recent_history.append(cast(Dict[str, Any], item))

    return {
        "metrics": metrics_snapshot,
        "vectorstore": {
            "docs_indexed": docs_indexed,
            "vectors_indexed": vectors_indexed,
        },
        "docs": docs,
        "docs_count": len(docs),
        "redis": redis_snapshot,
        "recent_statuses": recent_statuses[:20],
        "recent_history": recent_history[:40],
        "upload_limits": {
            "max_mb": settings.max_upload_mb,
            "allowed_exts": sorted(ALLOWED_DOC_EXTS),
        },
    }


@app.post("/api/admin/test_query", response_model=QueryResponse)
def admin_test_query(req: QueryRequest, admin: Tuple[str, str] = Depends(_require_admin)) -> QueryResponse:
    admin_user, _ = admin
    start = time.perf_counter()
    req_id = uuid.uuid4().hex[:8]
    logger.info("[admin-test:%s] start admin=%s top_k=%s", req_id, admin_user, req.top_k)
    try:
        resp, from_cache = orchestrator.handle(req)
        latency_ms = (time.perf_counter() - start) * 1000
        metrics.record_query(latency_ms=latency_ms, cached=from_cache, used_tools=len(resp.used_tools))
        logger.info("[admin-test:%s] done latency_ms=%.1f cached=%s", req_id, latency_ms, from_cache)
        return resp
    except FileNotFoundError as exc:
        metrics.record_error()
        logger.exception("[admin-test:%s] failed: %s", req_id, exc)
        raise HTTPException(status_code=400, detail=str(exc))


@app.post("/api/session/new", response_model=SessionCreateResponse)
def new_session(
    payload: Optional[SessionCreateRequest] = None,
    user: str = Depends(_require_user_active),
) -> SessionCreateResponse:
    title = (payload.title.strip() if payload and payload.title else None)
    summary = session_store.create_session(user, title=title)
    return SessionCreateResponse(
        session_id=summary.session_id,
        title=summary.title,
        created_at=summary.created_at,
    )


@app.get("/api/session", response_model=SessionListResponse)
def list_sessions(user: str = Depends(_require_user_active)) -> SessionListResponse:
    sessions = session_store.list_sessions(user, limit=50)
    return SessionListResponse(sessions=sessions)


@app.patch("/api/session/{session_id}", response_model=SessionSummary)
def rename_session_endpoint(
    session_id: str,
    payload: SessionUpdateRequest,
    user: str = Depends(_require_user_active),
) -> SessionSummary:
    summary = session_store.rename_session(session_id, user, payload.title)
    if not summary:
        raise HTTPException(status_code=404, detail="会话不存在或无权操作")
    return summary


@app.delete("/api/session/{session_id}")
def delete_session_endpoint(session_id: str, user: str = Depends(_require_user_active)) -> Dict[str, str]:
    ok = session_store.delete_session(session_id, user)
    if not ok:
        raise HTTPException(status_code=404, detail="会话不存在或无权删除")
    return {"status": "deleted"}


@app.get("/api/session/{session_id}/history", response_model=SessionHistoryResponse)
def session_history(session_id: str, user: str = Depends(_require_user_active)) -> SessionHistoryResponse:
    summary = session_store.ensure_session_for_user(session_id, user, create_if_missing=True)
    history = session_store.get_history(session_id)
    return SessionHistoryResponse(
        session_id=session_id,
        title=summary.title if summary else None,
        history=history,
    )


@app.get("/metrics")
def metrics_snapshot():
    return metrics.snapshot()


@app.get("/")
def root() -> FileResponse:
    index_path = Path(__file__).parent / "static" / "index.html"
    if not index_path.exists():
        raise HTTPException(status_code=404, detail="UI not found")
    return FileResponse(index_path)


@app.get("/chat")
def chat_page() -> FileResponse:
    chat_path = Path(__file__).parent / "static" / "chat.html"
    if not chat_path.exists():
        raise HTTPException(status_code=404, detail="Chat UI not found")
    return FileResponse(chat_path)


@app.get("/admin")
def admin_login_page() -> FileResponse:
    login_path = Path(__file__).parent / "static" / "admin_login.html"
    if not login_path.exists():
        raise HTTPException(status_code=404, detail="Admin login UI not found")
    return FileResponse(login_path)


@app.get("/admin/dashboard")
def admin_ui() -> FileResponse:
    admin_path = Path(__file__).parent / "static" / "admin.html"
    if not admin_path.exists():
        raise HTTPException(status_code=404, detail="Admin UI not found")
    return FileResponse(admin_path)


@app.get("/change-password")
def change_password_page() -> FileResponse:
    page_path = Path(__file__).parent / "static" / "change_password.html"
    if not page_path.exists():
        raise HTTPException(status_code=404, detail="Change password UI not found")
    return FileResponse(page_path)


def _persist_stream_history(req: QueryRequest, user: str, answer_parts: List[str], meta_payload: Dict[str, Any]):
    if not req.session_id:
        return
    _ = meta_payload  # reserved for future enrichment (e.g., sources)
    final_answer = "".join(answer_parts).strip()
    user_text = req.query or ""
    if not user_text and not final_answer:
        return
    try:
        session_store.ensure_session_for_user(req.session_id, user, create_if_missing=True)
        if user_text:
            session_store.add_message(req.session_id, "user", user_text, user=user)
        if final_answer:
            session_store.add_message(req.session_id, "assistant", final_answer, user=user)
    except Exception as exc:  # pragma: no cover - persistence failure shouldn't break stream
        logger.warning("[query-stream] failed to persist history for session=%s: %s", req.session_id, exc)
